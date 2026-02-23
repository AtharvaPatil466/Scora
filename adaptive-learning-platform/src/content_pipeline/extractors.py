# ==============================================================================
# FILE EXTRACTORS — One extractor per file type
# All return ExtractedContent for a unified downstream interface.
# ==============================================================================
from __future__ import annotations
import os
from dataclasses import dataclass, field, asdict
from typing import Optional


@dataclass
class ExtractedContent:
    """Unified output from any file extractor."""
    raw_text: str
    sections: list[dict] = field(default_factory=list)   # [{heading, body}]
    metadata: dict = field(default_factory=dict)          # filename, type, pages…
    images: list[str] = field(default_factory=list)       # paths to extracted imgs

    def to_dict(self) -> dict:
        return asdict(self)

    @classmethod
    def from_dict(cls, d: dict) -> "ExtractedContent":
        return cls(**d)


# ── PDF (PyMuPDF) ─────────────────────────────────────────────────────────────

def _extract_pdf(path: str) -> ExtractedContent:
    import fitz  # PyMuPDF

    doc = fitz.open(path)
    full_text_parts: list[str] = []
    sections: list[dict] = []

    for page_num, page in enumerate(doc, 1):
        text = page.get_text("text")
        full_text_parts.append(text)

        # Treat each page as a section
        lines = text.strip().split("\n")
        heading = lines[0].strip() if lines else f"Page {page_num}"
        body = "\n".join(lines[1:]).strip() if len(lines) > 1 else ""
        sections.append({"heading": heading, "body": body})

    doc.close()

    return ExtractedContent(
        raw_text="\n\n".join(full_text_parts),
        sections=sections,
        metadata={
            "filename": os.path.basename(path),
            "file_type": "pdf",
            "pages": len(sections),
        },
    )


# ── DOCX (python-docx) ───────────────────────────────────────────────────────

def _extract_docx(path: str) -> ExtractedContent:
    from docx import Document

    doc = Document(path)
    full_text_parts: list[str] = []
    sections: list[dict] = []
    current_heading: Optional[str] = None
    current_body_parts: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue

        full_text_parts.append(text)

        # If it's a heading style, start a new section
        if para.style and para.style.name and "Heading" in para.style.name:
            # Save previous section
            if current_heading is not None:
                sections.append({
                    "heading": current_heading,
                    "body": "\n".join(current_body_parts),
                })
            current_heading = text
            current_body_parts = []
        else:
            current_body_parts.append(text)

    # Save last section
    if current_heading is not None:
        sections.append({
            "heading": current_heading,
            "body": "\n".join(current_body_parts),
        })
    elif current_body_parts:
        sections.append({
            "heading": "Document Content",
            "body": "\n".join(current_body_parts),
        })

    return ExtractedContent(
        raw_text="\n\n".join(full_text_parts),
        sections=sections,
        metadata={
            "filename": os.path.basename(path),
            "file_type": "docx",
            "paragraphs": len(doc.paragraphs),
        },
    )


# ── PPTX (python-pptx) ───────────────────────────────────────────────────────

def _extract_pptx(path: str) -> ExtractedContent:
    from pptx import Presentation

    prs = Presentation(path)
    full_text_parts: list[str] = []
    sections: list[dict] = []

    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_texts.append(text)

        slide_text = "\n".join(slide_texts)
        full_text_parts.append(slide_text)

        heading = slide_texts[0] if slide_texts else f"Slide {slide_num}"
        body = "\n".join(slide_texts[1:]) if len(slide_texts) > 1 else ""
        sections.append({"heading": heading, "body": body})

        # Also check speaker notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                sections[-1]["body"] += f"\n\n[Speaker Notes] {notes}"
                full_text_parts.append(f"[Notes] {notes}")

    return ExtractedContent(
        raw_text="\n\n".join(full_text_parts),
        sections=sections,
        metadata={
            "filename": os.path.basename(path),
            "file_type": "pptx",
            "slides": len(prs.slides),
        },
    )


# ── Images / OCR (pytesseract + Pillow) ──────────────────────────────────────

def _extract_image(path: str) -> ExtractedContent:
    import pytesseract
    from PIL import Image

    img = Image.open(path)
    text = pytesseract.image_to_string(img)

    return ExtractedContent(
        raw_text=text.strip(),
        sections=[{"heading": "Extracted Text", "body": text.strip()}],
        metadata={
            "filename": os.path.basename(path),
            "file_type": "image",
            "dimensions": f"{img.width}x{img.height}",
        },
        images=[path],
    )


# ── Plain Text / Markdown ────────────────────────────────────────────────────

def _extract_text(path: str) -> ExtractedContent:
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        text = f.read()

    # For markdown, split on headings; for plain text, split on double newlines
    ext = os.path.splitext(path)[1].lower()
    sections: list[dict] = []

    if ext == ".md":
        import re
        parts = re.split(r"^(#{1,3}\s+.+)$", text, flags=re.MULTILINE)
        i = 0
        while i < len(parts):
            part = parts[i].strip()
            if part.startswith("#"):
                heading = part.lstrip("# ").strip()
                body = parts[i + 1].strip() if i + 1 < len(parts) else ""
                sections.append({"heading": heading, "body": body})
                i += 2
            else:
                if part:
                    sections.append({"heading": "Content", "body": part})
                i += 1
    else:
        # Plain text — split by double newlines into paragraphs
        paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
        if paragraphs:
            sections.append({
                "heading": paragraphs[0][:80],
                "body": "\n\n".join(paragraphs),
            })

    return ExtractedContent(
        raw_text=text,
        sections=sections or [{"heading": "Content", "body": text}],
        metadata={
            "filename": os.path.basename(path),
            "file_type": ext.lstrip(".") or "txt",
            "characters": len(text),
        },
    )


# ── Dispatcher ────────────────────────────────────────────────────────────────

_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
    ".png": _extract_image,
    ".jpg": _extract_image,
    ".jpeg": _extract_image,
    ".bmp": _extract_image,
    ".tiff": _extract_image,
    ".webp": _extract_image,
    ".txt": _extract_text,
    ".md": _extract_text,
}


def extract_file(path: str, original_filename: str | None = None) -> ExtractedContent:
    """
    Extract content from any supported file type.
    Falls back to plain-text extraction for unknown extensions.
    """
    name = original_filename or os.path.basename(path)
    ext = os.path.splitext(name)[1].lower()

    extractor = _EXTRACTORS.get(ext, _extract_text)
    return extractor(path)
